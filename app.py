"""
Email Marketing Dashboard — FastAPI Backend
============================================
Wraps the existing email-sending logic and exposes a REST + SSE API
consumed by the premium HTML dashboard (static/index.html).

Run:
    python app.py
    then open http://localhost:8000
"""

import asyncio
import csv
import imaplib
import io
import json
import os
import smtplib
import ssl
import tempfile
import threading
import time
from datetime import datetime
from email.mime.application import MIMEApplication
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from queue import Empty, Queue
from string import Formatter
from typing import Optional

import openpyxl
import uvicorn
from fastapi import FastAPI, File, Form, HTTPException, UploadFile, Request, Response
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, RedirectResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles

# ---------------------------------------------------------------------------
# Paths
# ---------------------------------------------------------------------------
ROOT = os.path.dirname(os.path.abspath(__file__))
TEMPLATES_DIR = os.path.join(ROOT, "templates")
LOG_PATH = os.path.join(ROOT, "sent_log.csv")
STATIC_DIR = os.path.join(ROOT, "static")

os.makedirs(STATIC_DIR, exist_ok=True)

# ---------------------------------------------------------------------------
# Platform presets
# ---------------------------------------------------------------------------
PLATFORM_PRESETS = {
    "gmail": {
        "smtp_host": "smtp.gmail.com",
        "smtp_port": 587,
        "smtp_security": "starttls",
        "imap_host": "imap.gmail.com",
        "imap_port": 993,
        "sent_folder": "[Gmail]/Sent Mail",
    },
    "hostinger": {
        "smtp_host": "smtp.hostinger.com",
        "smtp_port": 465,
        "smtp_security": "ssl",
        "imap_host": "imap.hostinger.com",
        "imap_port": 993,
        "sent_folder": "INBOX.Sent",
    },
    "outlook": {
        "smtp_host": "smtp.office365.com",
        "smtp_port": 587,
        "smtp_security": "starttls",
        "imap_host": "outlook.office365.com",
        "imap_port": 993,
        "sent_folder": "Sent Items",
    },
}

# ---------------------------------------------------------------------------
# Campaign state (single concurrent campaign)
# ---------------------------------------------------------------------------
_campaign_lock = threading.Lock()
_campaign_event_queue: Optional[Queue] = None
_campaign_stop_flag: Optional[threading.Event] = None
_campaign_thread: Optional[threading.Thread] = None

# ---------------------------------------------------------------------------
# App
# ---------------------------------------------------------------------------
app = FastAPI(title="Email Marketing Dashboard")

# Allow a separately-hosted frontend (e.g. on Vercel) to call this API.
# Set ALLOWED_ORIGINS="https://your-app.vercel.app" (comma-separated) in prod.
_origins = os.environ.get("ALLOWED_ORIGINS", "*").split(",")
app.add_middleware(
    CORSMiddleware,
    allow_origins=[o.strip() for o in _origins],
    allow_methods=["*"],
    allow_headers=["*"],
)

app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def placeholders(template: str) -> set:
    return {fn for _, fn, _, _ in Formatter().parse(template) if fn}


def render(template: str, row: dict) -> str:
    safe = {k: ("" if v is None else str(v)) for k, v in row.items()}
    for field in placeholders(template):
        safe.setdefault(field, "")
    return template.format(**safe)


def read_excel(data: bytes) -> tuple[list[str], list[dict]]:
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Excel file is empty.")
    headers = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    if "email" not in headers:
        raise ValueError("Excel must have an 'email' column.")

    def clean(val):
        if val is None:
            return ""
        return str(val).replace("\xa0", " ").strip()

    recipients = []
    seen = set()
    for raw in rows[1:]:
        row = {headers[i]: clean(raw[i]) for i in range(len(headers)) if headers[i]}
        email = "".join(row.get("email", "").split()).rstrip(",;")
        if not email or "@" not in email or " " in email:
            continue
        if email.lower() in seen:
            continue
        seen.add(email.lower())
        row["email"] = email
        name = " ".join(row.get("name", "").split())
        row["name"] = name if name else email.split("@")[0]
        recipients.append(row)
    return headers, recipients


def load_already_sent(campaign: str) -> set:
    sent = set()
    if not os.path.exists(LOG_PATH):
        return sent
    with open(LOG_PATH, newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            if r.get("campaign") == campaign and r.get("status") == "sent":
                sent.add(r.get("email", "").lower())
    return sent


def log_result(campaign: str, email: str, status: str, detail: str = ""):
    new = not os.path.exists(LOG_PATH)
    with open(LOG_PATH, "a", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        if new:
            w.writerow(["timestamp", "campaign", "email", "status", "detail"])
        w.writerow([datetime.now().isoformat(timespec="seconds"), campaign, email, status, detail])


def _smtp_candidates(host: str, port, security):
    """Primary (host, port, security) plus the sensible alternate port."""
    p = int(port)
    prim = (host, p, security or ("ssl" if p == 465 else "starttls"))
    cands = [prim]
    alt = (host, 587, "starttls") if p == 465 else (host, 465, "ssl")
    if alt != prim:
        cands.append(alt)
    return cands


def open_smtp_login(host, port, security, email, password, timeout=30, attempts=2):
    """
    Connect + log in to SMTP resiliently.
    - Generous timeout (some hosts' TLS handshake is slow, ~8s+).
    - Retries the primary port, then falls back to the alternate (465<->587).
    Returns (server, (host, port, security_used)). Raises the last error if all fail.
    """
    last = None
    for h, p, sec in _smtp_candidates(host, port, security):
        for _ in range(attempts):
            srv = None
            try:
                ctx = ssl.create_default_context()
                if sec == "ssl":
                    srv = smtplib.SMTP_SSL(h, p, context=ctx, timeout=timeout)
                else:
                    srv = smtplib.SMTP(h, p, timeout=timeout)
                    srv.ehlo()
                    srv.starttls(context=ctx)
                    srv.ehlo()
                srv.login(email, password)
                return srv, (h, p, sec)
            except Exception as e:  # noqa: BLE001
                last = e
                try:
                    if srv is not None:
                        srv.close()
                except Exception:  # noqa: BLE001
                    pass
    raise last if last else RuntimeError("SMTP connection failed.")


def build_message(sender_email: str, sender_name: str, to_email: str,
                  subject: str, html_body: str, attachments: list) -> MIMEMultipart:
    outer = MIMEMultipart("mixed") if attachments else MIMEMultipart("alternative")
    outer["From"] = f"{sender_name} <{sender_email}>"
    outer["To"] = to_email
    outer["Subject"] = subject

    if attachments:
        body_part = MIMEMultipart("alternative")
        body_part.attach(MIMEText(html_body, "html", "utf-8"))
        outer.attach(body_part)
        for path in attachments:
            full = path if os.path.isabs(path) else os.path.join(ROOT, path)
            if os.path.exists(full):
                with open(full, "rb") as f:
                    part = MIMEApplication(f.read(), _subtype="octet-stream")
                part.add_header("Content-Disposition", "attachment", filename=os.path.basename(full))
                outer.attach(part)
    else:
        outer.attach(MIMEText(html_body, "html", "utf-8"))
    return outer


# ---------------------------------------------------------------------------
# Campaign worker (runs in background thread)
# ---------------------------------------------------------------------------
def campaign_worker(cfg: dict, recipients: list, campaign: str,
                    subject_tpl: str, body_tpl: str,
                    delay_seconds: float, max_per_run: int,
                    event_q: Queue, stop_flag: threading.Event):
    def emit(kind: str, **kwargs):
        event_q.put({"type": kind, **kwargs})

    already = load_already_sent(campaign)
    queue = [r for r in recipients if r["email"].lower() not in already]
    queue = queue[:max_per_run]
    total = len(queue)

    if total == 0:
        emit("done", sent=0, failed=0, skipped=len(recipients), message="Nothing to send — all already sent.")
        return

    emit("start", total=total)

    # Open IMAP for Sent copies
    imap = None
    if cfg.get("save_to_sent", True):
        try:
            imap = imaplib.IMAP4_SSL(cfg["imap_host"], int(cfg["imap_port"]))
            imap.login(cfg["sender_email"], cfg["password"])
        except Exception as e:
            emit("warning", message=f"IMAP unavailable — Sent copies disabled: {e}")

    sent_count = failed_count = 0

    try:
        server, _used = open_smtp_login(
            cfg["smtp_host"], cfg["smtp_port"], cfg.get("smtp_security"),
            cfg["sender_email"], cfg["password"],
        )
    except Exception as e:
        emit("error", message=f"SMTP login failed: {e}")
        return

    with server:
        for i, r in enumerate(queue, 1):
            if stop_flag.is_set():
                emit("stopped", sent=sent_count, failed=failed_count)
                break

            subject = render(subject_tpl, r)
            body = render(body_tpl, r)
            msg = build_message(
                cfg["sender_email"], cfg.get("sender_name", cfg["sender_email"]),
                r["email"], subject, body, cfg.get("attachments", [])
            )
            try:
                server.sendmail(cfg["sender_email"], r["email"], msg.as_string())
                sent_count += 1
                log_result(campaign, r["email"], "sent")
                # Save to Sent folder via IMAP
                if imap:
                    try:
                        imap.append(
                            cfg.get("sent_folder", "Sent"),
                            "(\\Seen)",
                            imaplib.Time2Internaldate(time.time()),
                            msg.as_bytes()
                        )
                    except Exception:
                        pass
                emit("progress", index=i, total=total, email=r["email"],
                     status="sent", sent=sent_count, failed=failed_count)
            except Exception as e:
                failed_count += 1
                log_result(campaign, r["email"], "failed", str(e))
                emit("progress", index=i, total=total, email=r["email"],
                     status="failed", detail=str(e), sent=sent_count, failed=failed_count)

            if i < total and not stop_flag.is_set():
                time.sleep(delay_seconds)

    if imap:
        try:
            imap.logout()
        except Exception:
            pass

    if not stop_flag.is_set():
        emit("done", sent=sent_count, failed=failed_count, skipped=len(recipients) - total,
             message=f"Campaign complete. Sent: {sent_count}, Failed: {failed_count}")


# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------
AUTH_TOKEN = "mailflow_secure_session"

@app.post("/api/login")
async def login(payload: dict, response: Response):
    email = payload.get("email", "").strip()
    password = payload.get("password", "").strip()
    
    if email == "support@ankitkumaracademy.com" and password == "Kumar@20.26":
        response.set_cookie(key="auth_token", value=AUTH_TOKEN, httponly=True)
        return {"success": True}
    return {"success": False, "message": "Invalid credentials."}


@app.post("/api/logout")
async def logout(response: Response):
    # Clears the session cookie so the next load of "/" serves the landing page.
    response.delete_cookie(key="auth_token", path="/")
    return {"success": True}


@app.get("/logout")
async def logout_redirect():
    # Server-side logout: delete the cookie AND redirect home in one response.
    # Robust against browser caching / client-side timing issues.
    resp = RedirectResponse(url="/", status_code=303)
    resp.delete_cookie(key="auth_token", path="/")
    return resp


@app.get("/", response_class=HTMLResponse)
async def dashboard(request: Request):
    token = request.cookies.get("auth_token")
    if token == AUTH_TOKEN:
        html_path = os.path.join(STATIC_DIR, "index.html")
    else:
        html_path = os.path.join(STATIC_DIR, "landing.html")

    if not os.path.exists(html_path):
        return HTMLResponse("<h1>Page not found.</h1>", status_code=500)
    with open(html_path, encoding="utf-8") as f:
        # no-store so the browser never serves a stale (authed) page after logout.
        return HTMLResponse(f.read(), headers={"Cache-Control": "no-store, must-revalidate"})


@app.get("/config.js")
async def config_js():
    path = os.path.join(STATIC_DIR, "config.js")
    if os.path.exists(path):
        return FileResponse(path, media_type="application/javascript")
    return HTMLResponse("window.MAILFLOW_API_BASE='';", media_type="application/javascript")


@app.get("/api/presets")
async def get_presets():
    return PLATFORM_PRESETS


@app.get("/api/templates")
async def list_templates():
    if not os.path.isdir(TEMPLATES_DIR):
        return []
    result = []
    for name in sorted(os.listdir(TEMPLATES_DIR)):
        d = os.path.join(TEMPLATES_DIR, name)
        body_path = os.path.join(d, "body.html")
        subj_path = os.path.join(d, "subject.txt")
        if not os.path.isdir(d) or not os.path.exists(body_path):
            continue
        subject = "Hello {name}"
        if os.path.exists(subj_path):
            with open(subj_path, encoding="utf-8") as f:
                subject = f.read().strip()
        with open(body_path, encoding="utf-8") as f:
            body = f.read()
        fields = sorted(placeholders(subject) | placeholders(body))
        result.append({"name": name, "subject": subject, "body_preview": body[:300], "fields": fields})
    return result


def _safe_template_dir(name: str) -> str:
    """Resolve a template folder path, blocking path traversal."""
    safe = "".join(c for c in name if c.isalnum() or c in " -_").strip()
    safe = safe.replace(" ", "-").lower()
    if not safe:
        raise HTTPException(400, "Invalid template name.")
    d = os.path.abspath(os.path.join(TEMPLATES_DIR, safe))
    if not d.startswith(os.path.abspath(TEMPLATES_DIR) + os.sep):
        raise HTTPException(400, "Invalid template name.")
    return d


@app.get("/api/templates/{name}")
async def get_template(name: str):
    """Full template (subject + body + fields) for previewing/editing."""
    d = _safe_template_dir(name)
    body_path = os.path.join(d, "body.html")
    subj_path = os.path.join(d, "subject.txt")
    if not os.path.isdir(d) or not os.path.exists(body_path):
        raise HTTPException(404, "Template not found.")
    with open(body_path, encoding="utf-8") as f:
        body = f.read()
    subject = "Hello {name}"
    if os.path.exists(subj_path):
        with open(subj_path, encoding="utf-8") as f:
            subject = f.read().strip()
    fields = sorted(placeholders(subject) | placeholders(body))
    return {"name": os.path.basename(d), "subject": subject, "body": body, "fields": fields}


@app.get("/api/templates/{name}/preview", response_class=HTMLResponse)
async def preview_template(name: str):
    """Rendered HTML with sample data filled in, for an iframe preview."""
    d = _safe_template_dir(name)
    body_path = os.path.join(d, "body.html")
    if not os.path.isdir(d) or not os.path.exists(body_path):
        raise HTTPException(404, "Template not found.")
    with open(body_path, encoding="utf-8") as f:
        body = f.read()
    sample = {"name": "Alex Johnson", "email": "alex@example.com"}
    # Any other placeholders get a readable sample value like [company].
    for field in placeholders(body):
        sample.setdefault(field, f"[{field}]")
    return HTMLResponse(render(body, sample))


@app.post("/api/templates")
async def create_template(payload: dict):
    """Create/save a new template folder (persists to disk)."""
    name = payload.get("name", "").strip()
    subject = payload.get("subject", "").strip() or "Hello {name}"
    body = payload.get("body", "").strip()
    if not name:
        raise HTTPException(400, "Template name is required.")
    if not body:
        raise HTTPException(400, "Template body (HTML) is required.")

    d = _safe_template_dir(name)
    if os.path.exists(d) and not payload.get("overwrite"):
        raise HTTPException(409, "A template with that name already exists.")
    os.makedirs(d, exist_ok=True)
    with open(os.path.join(d, "subject.txt"), "w", encoding="utf-8") as f:
        f.write(subject)
    with open(os.path.join(d, "body.html"), "w", encoding="utf-8") as f:
        f.write(body)

    fields = sorted(placeholders(subject) | placeholders(body))
    return {"name": os.path.basename(d), "subject": subject, "fields": fields}


@app.post("/api/test-connection")
async def test_connection(payload: dict):
    email = payload.get("email", "").strip()
    password = payload.get("password", "").strip()
    platform = payload.get("platform", "hostinger")
    preset = PLATFORM_PRESETS.get(platform, PLATFORM_PRESETS["hostinger"])

    results = {}

    # Test SMTP (resilient: generous timeout, retry, port fallback)
    try:
        srv, used = open_smtp_login(
            preset["smtp_host"], preset["smtp_port"], preset["smtp_security"],
            email, password,
        )
        try:
            srv.quit()
        except Exception:  # noqa: BLE001
            pass
        results["smtp"] = "ok"
        results["smtp_via"] = f"{used[0]}:{used[1]} ({used[2]})"
    except Exception as e:
        results["smtp"] = f"failed: {e}"

    # Test IMAP
    try:
        imap = imaplib.IMAP4_SSL(preset["imap_host"], int(preset["imap_port"]))
        imap.login(email, password)
        imap.logout()
        results["imap"] = "ok"
    except Exception as e:
        results["imap"] = f"failed: {e}"

    success = results["smtp"] == "ok" and results["imap"] == "ok"
    return {"success": success, "smtp": results["smtp"], "imap": results["imap"]}


@app.post("/api/upload-excel")
async def upload_excel(file: UploadFile = File(...)):
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(400, "Only .xlsx / .xls files are supported.")
    data = await file.read()
    try:
        headers, recipients = read_excel(data)
    except ValueError as e:
        raise HTTPException(400, str(e))

    preview = recipients[:5]
    return {
        "total": len(recipients),
        "headers": [h for h in headers if h],
        "preview": preview,
        "filename": file.filename,
    }


@app.post("/api/start-campaign")
async def start_campaign(payload: dict):
    global _campaign_event_queue, _campaign_stop_flag, _campaign_thread

    with _campaign_lock:
        if _campaign_thread and _campaign_thread.is_alive():
            raise HTTPException(409, "A campaign is already running. Stop it first.")

        platform = payload.get("platform", "hostinger")
        preset = PLATFORM_PRESETS.get(platform, PLATFORM_PRESETS["hostinger"])

        cfg = {
            "sender_email": payload["email"],
            "sender_name": payload.get("sender_name", payload["email"]),
            "password": payload["password"],
            **preset,
            "save_to_sent": True,
            "attachments": [],
        }

        excel_data = bytes(payload["excel_data"])  # base64 decoded bytes passed as int list
        campaign = payload["campaign"]
        delay = float(payload.get("delay_seconds", 5))
        max_run = int(payload.get("max_per_run", 100))

        # Load template
        subj_path = os.path.join(TEMPLATES_DIR, campaign, "subject.txt")
        body_path = os.path.join(TEMPLATES_DIR, campaign, "body.html")
        if not os.path.exists(body_path):
            raise HTTPException(400, f"Template '{campaign}' not found.")
        with open(body_path, encoding="utf-8") as f:
            body_tpl = f.read()
        subject_tpl = "Hello {name}"
        if os.path.exists(subj_path):
            with open(subj_path, encoding="utf-8") as f:
                subject_tpl = f.read().strip()

        _, recipients = read_excel(excel_data)

        _campaign_event_queue = Queue()
        _campaign_stop_flag = threading.Event()
        _campaign_thread = threading.Thread(
            target=campaign_worker,
            args=(cfg, recipients, campaign, subject_tpl, body_tpl,
                  delay, max_run, _campaign_event_queue, _campaign_stop_flag),
            daemon=True,
        )
        _campaign_thread.start()

    return {"status": "started"}


@app.post("/api/stop-campaign")
async def stop_campaign():
    global _campaign_stop_flag
    if _campaign_stop_flag:
        _campaign_stop_flag.set()
    return {"status": "stopping"}


@app.get("/api/campaign-events")
async def campaign_events():
    """SSE stream of campaign progress events."""
    global _campaign_event_queue

    if _campaign_event_queue is None:
        raise HTTPException(400, "No campaign running.")

    q = _campaign_event_queue

    async def event_generator():
        while True:
            try:
                event = q.get(timeout=0.5)
                yield f"data: {json.dumps(event)}\n\n"
                if event["type"] in ("done", "error", "stopped"):
                    break
            except Empty:
                yield ": ping\n\n"
                await asyncio.sleep(0.1)

    return StreamingResponse(event_generator(), media_type="text/event-stream",
                             headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@app.get("/api/logs")
async def get_logs(limit: int = 50):
    if not os.path.exists(LOG_PATH):
        return []
    rows = []
    with open(LOG_PATH, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            rows.append(row)
    return rows[-limit:]


@app.get("/api/status")
async def status():
    running = bool(_campaign_thread and _campaign_thread.is_alive())
    return {"running": running}


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------
if __name__ == "__main__":
    port = int(os.environ.get("PORT", 8000))
    # Only auto-open a browser when running locally (not on a cloud host).
    if not os.environ.get("PORT"):
        import webbrowser
        print("\n  Email Marketing Dashboard")
        print(f"  http://localhost:{port}\n")
        webbrowser.open(f"http://localhost:{port}")
    uvicorn.run("app:app", host="0.0.0.0", port=port, reload=False)
